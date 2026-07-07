from pathlib import Path
import math
import json

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import requests
import seaborn as sns
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


BASE = Path(r"e:\Aneurysms Genomics Project")
RES = BASE / "results"
MOYA = RES / "moyamoya"
OUT = RES / "comparative_analysis"
FIG = OUT / "figures_ultra_hd"
OUT.mkdir(parents=True, exist_ok=True)
FIG.mkdir(parents=True, exist_ok=True)


def fetch_symbol_map(entrez_ids):
    ids = sorted({str(int(float(x))) for x in entrez_ids if pd.notna(x) and str(x).strip()})
    out = {}
    if not ids:
        return out
    step = 800
    for i in range(0, len(ids), step):
        batch = ids[i:i + step]
        try:
            r = requests.post(
                "https://mygene.info/v3/gene",
                data={"ids": ",".join(batch), "fields": "symbol,name,entrezgene"},
                timeout=60,
            )
            if r.status_code != 200:
                continue
            data = r.json()
            if isinstance(data, dict):
                data = [data]
            for row in data:
                q = str(row.get("query", "")).strip()
                sym = row.get("symbol", None)
                if q and sym:
                    out[q] = sym
        except Exception:
            continue
    return out


def load_aneurysm():
    d = pd.read_csv(RES / "gse122897_deseq2_results.tsv", sep="\t")
    d["Gene_ID"] = d["Gene_ID"].astype(str).str.replace(".0", "", regex=False)
    sym_map = fetch_symbol_map(d["Gene_ID"].tolist())
    d["GeneSymbol"] = d["Gene_ID"].map(sym_map)
    d = d.dropna(subset=["GeneSymbol"]).copy()
    d = d.sort_values("padj")
    d = d.drop_duplicates("GeneSymbol", keep="first")
    d["Direction"] = np.where(d["log2FoldChange"] >= 0, "Up_in_ruptured", "Down_in_ruptured")
    sig = d[(d["padj"] < 0.05) & (np.abs(d["log2FoldChange"]) >= 1)].copy()
    return d, sig


def load_moya():
    mia = pd.read_csv(MOYA / "gse157628_mmd_vs_ia_de.tsv", sep="\t")
    mepi = pd.read_csv(MOYA / "gse157628_mmd_vs_epi_de.tsv", sep="\t")
    for df in (mia, mepi):
        df["GeneSymbol"] = df["GeneSymbol"].astype(str).str.strip()
        df.loc[df["GeneSymbol"].isin(["", "nan", "None"]), "GeneSymbol"] = np.nan
        df.dropna(subset=["GeneSymbol"], inplace=True)
        df.sort_values("P_value", inplace=True)
        df.drop_duplicates("GeneSymbol", keep="first", inplace=True)
    mia["Direction"] = np.where(mia["Log2FC_case_vs_ctrl"] >= 0, "Up_in_MMD", "Down_in_MMD")
    mepi["Direction"] = np.where(mepi["Log2FC_case_vs_ctrl"] >= 0, "Up_in_MMD", "Down_in_MMD")

    mia_focus = mia[(mia["P_value"] < 1e-3) & (np.abs(mia["Log2FC_case_vs_ctrl"]) >= 1)].copy()
    mepi_focus = mepi[(mepi["FDR_BH"] < 0.05) & (np.abs(mepi["Log2FC_case_vs_ctrl"]) >= 1)].copy()
    return mia, mepi, mia_focus, mepi_focus


def draw_three_set_overlap(ax, A, B, C, labels):
    # Hand-drawn 3-set diagram with circles and annotated counts.
    circles = [
        ((-0.8, 0.1), 1.25, "#4c72b0"),
        ((0.8, 0.1), 1.25, "#dd8452"),
        ((0.0, -0.9), 1.25, "#55a868"),
    ]
    for (xy, r, c) in circles:
        patch = plt.Circle(xy, r, color=c, alpha=0.25, ec=c, lw=2)
        ax.add_patch(patch)

    only_a = len(A - B - C)
    only_b = len(B - A - C)
    only_c = len(C - A - B)
    ab = len((A & B) - C)
    ac = len((A & C) - B)
    bc = len((B & C) - A)
    abc = len(A & B & C)

    ax.text(-1.55, 0.35, str(only_a), fontsize=16, fontweight="bold")
    ax.text(1.45, 0.35, str(only_b), fontsize=16, fontweight="bold")
    ax.text(0.0, -1.95, str(only_c), fontsize=16, fontweight="bold", ha="center")
    ax.text(0.0, 0.45, str(ab), fontsize=16, fontweight="bold", ha="center")
    ax.text(-0.58, -0.65, str(ac), fontsize=16, fontweight="bold", ha="center")
    ax.text(0.58, -0.65, str(bc), fontsize=16, fontweight="bold", ha="center")
    ax.text(0.0, -0.15, str(abc), fontsize=18, fontweight="bold", ha="center", color="black")

    ax.text(-1.7, 1.45, labels[0], fontsize=12, fontweight="bold", color="#4c72b0")
    ax.text(0.9, 1.45, labels[1], fontsize=12, fontweight="bold", color="#dd8452")
    ax.text(-0.5, -2.45, labels[2], fontsize=12, fontweight="bold", color="#55a868")
    ax.set_xlim(-2.2, 2.2)
    ax.set_ylim(-2.7, 2.0)
    ax.axis("off")


def make_overlap_figure(set_aneurysm, set_mia, set_mepi):
    plt.figure(figsize=(14, 10), dpi=500)
    ax = plt.gca()
    draw_three_set_overlap(
        ax,
        set_aneurysm,
        set_mia,
        set_mepi,
        [
            "Aneurysm (Ruptured vs Unruptured)\nDESeq2 FDR<0.05 |log2FC|>=1",
            "Moya Moya (MMD vs IA)\nP<0.001 |log2FC|>=1",
            "Moya Moya (MMD vs EPI)\nFDR<0.05 |log2FC|>=1",
        ],
    )
    plt.title("Cross-condition gene overlap landscape", fontsize=20, fontweight="bold")
    out = FIG / "gene_overlap_3set_ultra_hd.png"
    plt.tight_layout()
    plt.savefig(out, dpi=600)
    plt.close()
    return out


def make_effect_scatter(aneurysm_all, mia_all, mepi_all):
    merged_mia = aneurysm_all[["GeneSymbol", "log2FoldChange"]].merge(
        mia_all[["GeneSymbol", "Log2FC_case_vs_ctrl", "P_value"]],
        on="GeneSymbol",
        how="inner",
    )
    merged_mepi = aneurysm_all[["GeneSymbol", "log2FoldChange"]].merge(
        mepi_all[["GeneSymbol", "Log2FC_case_vs_ctrl", "FDR_BH"]],
        on="GeneSymbol",
        how="inner",
    )

    fig, axes = plt.subplots(1, 2, figsize=(20, 9), dpi=500)
    sns.set_style("whitegrid")

    x = merged_mia["log2FoldChange"].values
    y = merged_mia["Log2FC_case_vs_ctrl"].values
    axes[0].scatter(x, y, s=30, alpha=0.5, color="#4c72b0")
    if len(merged_mia) > 1:
        m, b = np.polyfit(x, y, 1)
        xx = np.linspace(x.min(), x.max(), 100)
        axes[0].plot(xx, m * xx + b, color="black", lw=2)
    axes[0].axhline(0, ls="--", c="gray")
    axes[0].axvline(0, ls="--", c="gray")
    axes[0].set_title("Aneurysm vs MMD-IA effect concordance", fontweight="bold")
    axes[0].set_xlabel("Aneurysm log2FC (ruptured vs unruptured)")
    axes[0].set_ylabel("Moya Moya log2FC (MMD vs IA)")

    x2 = merged_mepi["log2FoldChange"].values
    y2 = merged_mepi["Log2FC_case_vs_ctrl"].values
    axes[1].scatter(x2, y2, s=30, alpha=0.5, color="#dd8452")
    if len(merged_mepi) > 1:
        m2, b2 = np.polyfit(x2, y2, 1)
        xx2 = np.linspace(x2.min(), x2.max(), 100)
        axes[1].plot(xx2, m2 * xx2 + b2, color="black", lw=2)
    axes[1].axhline(0, ls="--", c="gray")
    axes[1].axvline(0, ls="--", c="gray")
    axes[1].set_title("Aneurysm vs MMD-EPI effect concordance", fontweight="bold")
    axes[1].set_xlabel("Aneurysm log2FC (ruptured vs unruptured)")
    axes[1].set_ylabel("Moya Moya log2FC (MMD vs EPI)")

    out = FIG / "effect_concordance_scatter_ultra_hd.png"
    plt.suptitle("Cross-disease effect-size geometry (shared gene symbols)", fontsize=20, fontweight="bold")
    plt.tight_layout(rect=[0, 0, 1, 0.95])
    plt.savefig(out, dpi=600)
    plt.close()
    return out, merged_mia, merged_mepi


def make_shared_heatmap(aneurysm_all, mia_all, mepi_all):
    m = aneurysm_all[["GeneSymbol", "log2FoldChange", "padj"]].merge(
        mia_all[["GeneSymbol", "Log2FC_case_vs_ctrl", "P_value"]],
        on="GeneSymbol",
        how="inner",
    ).merge(
        mepi_all[["GeneSymbol", "Log2FC_case_vs_ctrl", "FDR_BH"]],
        on="GeneSymbol",
        how="inner",
        suffixes=("_MIA", "_MEPI"),
    )
    if len(m) == 0:
        return None, m
    m["score"] = (
        np.abs(m["log2FoldChange"]) +
        np.abs(m["Log2FC_case_vs_ctrl_MIA"]) +
        np.abs(m["Log2FC_case_vs_ctrl_MEPI"])
    )
    top = m.sort_values("score", ascending=False).head(50).copy()
    mat = top.set_index("GeneSymbol")[["log2FoldChange", "Log2FC_case_vs_ctrl_MIA", "Log2FC_case_vs_ctrl_MEPI"]]
    mat.columns = ["Aneurysm_RvU", "Moya_MMDvIA", "Moya_MMDvEPI"]
    z = (mat - mat.mean(axis=0)) / mat.std(axis=0).replace(0, 1)

    cg = sns.clustermap(
        z,
        cmap="vlag",
        center=0,
        linewidths=0.1,
        figsize=(12, 18),
        dendrogram_ratio=(0.15, 0.08),
        cbar_kws={"label": "Z-scored effect size"},
    )
    cg.fig.suptitle("Top shared genes: clustered effect signature map", y=1.02, fontsize=18, fontweight="bold")
    out = FIG / "shared_gene_clustermap_ultra_hd.png"
    cg.fig.savefig(out, dpi=600, bbox_inches="tight")
    plt.close(cg.fig)
    return out, top


def pathway_theme_counts():
    a = pd.read_csv(RES / "gse122897_pathway_enrichment_filtered.tsv", sep="\t")
    mepi_file = MOYA / "gse157628_mmd_vs_epi_pathway_strict_ranked.tsv"
    mia_file = MOYA / "gse157628_mmd_vs_ia_pathway_strict_ranked.tsv"
    mepi = pd.read_csv(mepi_file, sep="\t") if mepi_file.exists() else pd.DataFrame(columns=["Term_Name"])
    mia = pd.read_csv(mia_file, sep="\t") if mia_file.exists() else pd.DataFrame(columns=["Term_Name"])

    themes = {
        "Inflammation": ["interleukin", "cytokine", "immune", "leukocyte", "chemotaxis", "inflamm"],
        "ECM_Remodeling": ["extracellular matrix", "collagen", "matrix", "integrin", "adhesion"],
        "Vascular_Motility": ["migration", "locomotion", "motility", "vasculature", "blood vessel"],
        "Oxidative_Stress": ["oxid", "hypoxia", "reactive oxygen", "stress"],
        "Signal_Transduction": ["nf", "mapk", "receptor", "signaling"],
    }

    def count_theme(df, kws):
        names = df["Term_Name"].astype(str).str.lower()
        return int(sum(names.str.contains("|".join(kws), regex=True)))

    rows = []
    for theme, kws in themes.items():
        rows.append(
            {
                "Theme": theme,
                "Aneurysm": count_theme(a, kws),
                "MMD_vs_IA": count_theme(mia, kws),
                "MMD_vs_EPI": count_theme(mepi, kws),
            }
        )
    return pd.DataFrame(rows)


def make_theme_radar(theme_df):
    labels = theme_df["Theme"].tolist()
    categories = len(labels)
    angles = np.linspace(0, 2 * np.pi, categories, endpoint=False).tolist()
    angles += angles[:1]

    fig = plt.figure(figsize=(11, 11), dpi=500)
    ax = plt.subplot(111, polar=True)

    for col, color in [("Aneurysm", "#4c72b0"), ("MMD_vs_IA", "#55a868"), ("MMD_vs_EPI", "#dd8452")]:
        vals = theme_df[col].tolist()
        vals += vals[:1]
        ax.plot(angles, vals, linewidth=2.5, label=col, color=color)
        ax.fill(angles, vals, alpha=0.1, color=color)

    ax.set_xticks(angles[:-1])
    ax.set_xticklabels(labels, fontsize=12)
    ax.set_title("Biological pathway-theme footprint across conditions", fontsize=18, fontweight="bold", pad=20)
    ax.legend(loc="upper right", bbox_to_anchor=(1.2, 1.12))
    out = FIG / "pathway_theme_radar_ultra_hd.png"
    plt.tight_layout()
    plt.savefig(out, dpi=600)
    plt.close()
    return out


def top_gene_tables(aneurysm_sig, mia_focus, mepi_focus):
    a = aneurysm_sig[["GeneSymbol", "log2FoldChange", "padj"]].copy().head(25)
    a.columns = ["GeneSymbol", "Aneurysm_log2FC", "Aneurysm_padj"]

    mia = mia_focus[["GeneSymbol", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH"]].copy().head(25)
    mia.columns = ["GeneSymbol", "MMDvsIA_log2FC", "MMDvsIA_P", "MMDvsIA_FDR"]

    mepi = mepi_focus[["GeneSymbol", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH"]].copy().head(25)
    mepi.columns = ["GeneSymbol", "MMDvsEPI_log2FC", "MMDvsEPI_P", "MMDvsEPI_FDR"]
    return a, mia, mepi


def add_table(doc, df, title):
    doc.add_heading(title, 2)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        table.rows[0].cells[i].text = str(c)
    for _, r in df.iterrows():
        row = table.add_row().cells
        for i, v in enumerate(r):
            if isinstance(v, float):
                row[i].text = f"{v:.4g}"
            else:
                row[i].text = str(v)
    doc.add_paragraph("")


def build_doc(fig_paths, a_tbl, mia_tbl, mepi_tbl, overlap_stats, comp_stats):
    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name = "Calibri"
    normal.size = Pt(11)

    title = doc.add_heading("Comparative Transcriptomic Report: Aneurysm vs Moya Moya", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Focus: cross-condition gene signatures, effect-direction concordance, and high-resolution integrative visualization."
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    doc.add_heading("1. Key Comparative Findings", 1)
    doc.add_paragraph(
        f"Aneurysm robust DE signal: {comp_stats['aneurysm_sig_genes']} genes (DESeq2 padj<0.05 & |log2FC|>=1). "
        f"MMD vs IA focused set: {comp_stats['mmd_ia_focus_genes']} genes (P<0.001 & |log2FC|>=1). "
        f"MMD vs EPI focused set: {comp_stats['mmd_epi_focus_genes']} genes (FDR<0.05 & |log2FC|>=1)."
    )
    doc.add_paragraph(
        f"Shared genes: aneurysm ∩ MMDvsIA = {overlap_stats['a_mia']}, "
        f"aneurysm ∩ MMDvsEPI = {overlap_stats['a_mepi']}, "
        f"three-way overlap = {overlap_stats['a_mia_mepi']}."
    )
    doc.add_paragraph(
        "Interpretation: aneurysm findings are broad and pathway-rich, while Moya Moya shows a weak age-proximal contrast (MMDvsIA) and a stronger but confound-prone contrast (MMDvsEPI)."
    )

    doc.add_heading("2. Ultra-HD Comparative Illustrations", 1)
    for caption, p in fig_paths:
        if p is not None and Path(p).exists():
            doc.add_picture(str(p), width=Inches(7.2))
            cp = doc.add_paragraph(caption)
            cp.runs[0].italic = True
            doc.add_paragraph("")

    add_table(doc, a_tbl, "3. Top Aneurysm Genes (Ruptured vs Unruptured)")
    add_table(doc, mia_tbl, "4. Top Moya Moya Genes (MMD vs IA)")
    add_table(doc, mepi_tbl, "5. Top Moya Moya Genes (MMD vs EPI)")

    doc.add_heading("6. High-level Contrast Summary", 1)
    doc.add_paragraph(
        "- Aneurysm: strong, coherent transcriptomic shift with extensive enriched pathways.\n"
        "- MMD vs IA: low-powered signal with limited robust significance after strict correction.\n"
        "- MMD vs EPI: stronger differences but heavily susceptible to cohort-structure confounding.\n"
        "- Shared genes provide hypothesis-level links but require validation in balanced cohorts."
    )

    out_doc = OUT / "Aneurysm_vs_MoyaMoya_Comparative_Advanced_Report.docx"
    doc.save(out_doc)
    return out_doc


def main():
    aneurysm_all, aneurysm_sig = load_aneurysm()
    mia_all, mepi_all, mia_focus, mepi_focus = load_moya()

    set_a = set(aneurysm_sig["GeneSymbol"].tolist())
    set_mia = set(mia_focus["GeneSymbol"].tolist())
    set_mepi = set(mepi_focus["GeneSymbol"].tolist())

    f1 = make_overlap_figure(set_a, set_mia, set_mepi)
    f2, merged_mia, merged_mepi = make_effect_scatter(aneurysm_all, mia_all, mepi_all)
    f3, top_shared = make_shared_heatmap(aneurysm_all, mia_all, mepi_all)
    theme_df = pathway_theme_counts()
    f4 = make_theme_radar(theme_df)

    a_tbl, mia_tbl, mepi_tbl = top_gene_tables(aneurysm_sig, mia_focus, mepi_focus)
    a_tbl.to_csv(OUT / "top_aneurysm_genes_for_comparison.tsv", sep="\t", index=False)
    mia_tbl.to_csv(OUT / "top_mmd_vs_ia_genes_for_comparison.tsv", sep="\t", index=False)
    mepi_tbl.to_csv(OUT / "top_mmd_vs_epi_genes_for_comparison.tsv", sep="\t", index=False)
    merged_mia.to_csv(OUT / "shared_genes_aneurysm_vs_mmd_ia.tsv", sep="\t", index=False)
    merged_mepi.to_csv(OUT / "shared_genes_aneurysm_vs_mmd_epi.tsv", sep="\t", index=False)

    # Strict-overlap tables: genes significant in BOTH aneurysm and Moya Moya
    strict_mia = aneurysm_sig.merge(
        mia_focus[["GeneSymbol", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH"]],
        on="GeneSymbol",
        how="inner",
        suffixes=("", "_MIA"),
    )
    strict_mepi = aneurysm_sig.merge(
        mepi_focus[["GeneSymbol", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH"]],
        on="GeneSymbol",
        how="inner",
        suffixes=("", "_MEPI"),
    )
    strict_mia.to_csv(OUT / "strict_overlap_aneurysm_mmd_ia.tsv", sep="\t", index=False)
    strict_mepi.to_csv(OUT / "strict_overlap_aneurysm_mmd_epi.tsv", sep="\t", index=False)
    triple_overlap = aneurysm_sig.merge(
        mia_focus[["GeneSymbol", "Log2FC_case_vs_ctrl"]].rename(columns={"Log2FC_case_vs_ctrl": "Log2FC_MIA"}),
        on="GeneSymbol",
        how="inner",
    ).merge(
        mepi_focus[["GeneSymbol", "Log2FC_case_vs_ctrl"]].rename(columns={"Log2FC_case_vs_ctrl": "Log2FC_MEPI"}),
        on="GeneSymbol",
        how="inner",
    )
    triple_overlap.to_csv(OUT / "strict_overlap_triple.tsv", sep="\t", index=False)

    if f3 is not None and top_shared is not None:
        top_shared.to_csv(OUT / "top_shared_genes_effect_matrix.tsv", sep="\t", index=False)
    theme_df.to_csv(OUT / "pathway_theme_keyword_counts.tsv", sep="\t", index=False)

    overlap_stats = {
        "a_mia": len(set_a & set_mia),
        "a_mepi": len(set_a & set_mepi),
        "a_mia_mepi": len(set_a & set_mia & set_mepi),
    }
    comp_stats = {
        "aneurysm_sig_genes": len(set_a),
        "mmd_ia_focus_genes": len(set_mia),
        "mmd_epi_focus_genes": len(set_mepi),
    }

    fig_paths = [
        ("Figure 1. Three-set overlap map of cross-condition gene signatures.", f1),
        ("Figure 2. Effect-size concordance scatter between aneurysm and Moya Moya contrasts.", f2),
        ("Figure 3. Clustered heatmap of top shared genes across all contrasts.", f3),
        ("Figure 4. Pathway-theme radar footprint using enrichment keyword themes.", f4),
    ]
    out_doc = build_doc(fig_paths, a_tbl, mia_tbl, mepi_tbl, overlap_stats, comp_stats)

    summary = {
        "output_doc": str(out_doc),
        "figure_dir": str(FIG),
        "overlap_stats": overlap_stats,
        "comparison_stats": comp_stats,
    }
    with open(OUT / "comparative_report_summary.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
