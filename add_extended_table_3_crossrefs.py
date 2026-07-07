"""
Add cross-references to Extended Data Table 3 in the main manuscript and
Supplementary_Tables_Revised.docx.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

BASE = Path(r"e:\Aneurysms Genomics Project")
SUB = BASE / "Latest Submission Files"
MANUSCRIPT = SUB / "NatMed_Analysis_Final_Revised.docx"
SUPP_TABLES = SUB / "Supplementary_Tables_Revised.docx"
EDT3_FILE = "Extended_Data_Table_3_300gene_shortlist.tsv"
MARKER = "Extended Data Table 3"


def _replace_in_paragraph(para, old: str, new: str) -> bool:
    if old not in para.text:
        return False
    para.text = para.text.replace(old, new, 1)
    return True


def _patch_paragraphs(doc: Document, replacements: list[tuple[str, str]]) -> list[str]:
    applied: list[str] = []
    for para in doc.paragraphs:
        for old, new in replacements:
            if old in para.text and new not in para.text:
                if _replace_in_paragraph(para, old, new):
                    applied.append(old[:60] + ("..." if len(old) > 60 else ""))
    return applied


def patch_manuscript() -> None:
    if not MANUSCRIPT.exists():
        raise FileNotFoundError(MANUSCRIPT)

    doc = Document(str(MANUSCRIPT))
    if any(MARKER in p.text for p in doc.paragraphs):
        print(f"{MANUSCRIPT.name}: cross-references already present; skipping.")
        return

    replacements = [
        (
            "The full 300-gene shortlist is provided in Supplementary Table S2.",
            (
                f"The full 300-gene shortlist is provided in Extended Data Table 3 "
                f"(supplementary data file {EDT3_FILE}); the top 40 ranked genes are "
                "shown in Supplementary Table S2."
            ),
        ),
        (
            "(full shortlist in Supplementary Table S2).",
            (
                f"(complete ranked shortlist in Extended Data Table 3 and supplementary "
                f"data file {EDT3_FILE}; top 40 rows in Supplementary Table S2)."
            ),
        ),
        (
            "Background N = 300 is the most scientifically defensible choice; alternative backgrounds are reported in Supplementary Table S12.",
            (
                "Background N = 300 is the most scientifically defensible choice "
                f"(pre-specified shortlist; Extended Data Table 3); alternative backgrounds "
                "are reported in Supplementary Table S12."
            ),
        ),
    ]

    applied = _patch_paragraphs(doc, replacements)
    if not applied:
        raise RuntimeError(f"No manuscript paragraphs matched in {MANUSCRIPT.name}")

    doc.save(str(MANUSCRIPT))
    print(f"Updated {MANUSCRIPT.name}: {len(applied)} replacement(s)")


def patch_supplementary_tables() -> None:
    if not SUPP_TABLES.exists():
        raise FileNotFoundError(SUPP_TABLES)

    doc = Document(str(SUPP_TABLES))

    # Idempotent: only skip if S2 legend already names Extended Data Table 3
    s2_legend_done = any(
        p.text.strip().startswith("All 300 genes met") and MARKER in p.text
        for p in doc.paragraphs
    )
    if s2_legend_done:
        print(f"{SUPP_TABLES.name}: Table S2 cross-reference already present.")
    else:
        replacements = [
            (
                "Full 300-gene table in supplementary data file.",
                (
                    f"Full 300-gene ranked shortlist: Extended Data Table 3 "
                    f"(supplementary data file {EDT3_FILE})."
                ),
            ),
            (
                "N = 300 is the most scientifically defensible choice as both analyses drew from this pre-specified shortlist.",
                (
                    "N = 300 is the most scientifically defensible choice as both analyses "
                    f"drew from this pre-specified shortlist (Extended Data Table 3; "
                    f"{EDT3_FILE})."
                ),
            ),
        ]
        applied = _patch_paragraphs(doc, replacements)
        if not applied:
            raise RuntimeError(f"No supplementary-table paragraphs matched in {SUPP_TABLES.name}")
        print(f"Updated {SUPP_TABLES.name}: {len(applied)} replacement(s)")

    # Footnote on S2: add EDT3 pointer if missing
    for para in doc.paragraphs:
        if para.text.strip().startswith("Footnote (effect-size definition)") and MARKER not in para.text:
            if not para.text.rstrip().endswith("."):
                para.text = para.text.rstrip() + "."
            para.text = (
                para.text.rstrip()
                + f" The complete 300-gene shortlist is in Extended Data Table 3 "
                f"({EDT3_FILE})."
            )
            doc.save(str(SUPP_TABLES))
            print(f"Updated {SUPP_TABLES.name}: S2 footnote cross-reference")
            return

    doc.save(str(SUPP_TABLES))


def main() -> None:
    patch_manuscript()
    patch_supplementary_tables()
    print("Done.")


if __name__ == "__main__":
    main()
