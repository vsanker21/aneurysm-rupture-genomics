from pathlib import Path
import json

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


BASE = Path(r"e:\Aneurysms Genomics Project")
RES = BASE / "results"
CMP = RES / "comparative_analysis"
FIG = CMP / "figures_ultra_hd"


def add_figure(doc, path, caption):
    p = Path(path)
    if p.exists():
        doc.add_picture(str(p), width=Inches(7.0))
        cap = doc.add_paragraph(caption)
        cap.runs[0].italic = True
        doc.add_paragraph("")


def add_legend_paragraph(doc, heading, text):
    p = doc.add_paragraph()
    p.add_run(f"{heading}: ").bold = True
    p.add_run(text)


def add_table(doc, df, title):
    doc.add_heading(title, 3)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
    for _, r in df.iterrows():
        row = t.add_row().cells
        for i, v in enumerate(r):
            if isinstance(v, float):
                row[i].text = f"{v:.4g}"
            else:
                row[i].text = str(v)
    doc.add_paragraph("")


def main():
    summary = json.loads((CMP / "comparative_report_summary.json").read_text(encoding="utf-8"))
    moya = json.loads((RES / "moyamoya" / "gse157628_extended_summary.json").read_text(encoding="utf-8"))
    nested = pd.read_csv(RES / "gse122897_nested_de_cv_summary.tsv", sep="\t")
    top_aneurysm = pd.read_csv(CMP / "top_aneurysm_genes_for_comparison.tsv", sep="\t").head(15)
    top_mia = pd.read_csv(CMP / "top_mmd_vs_ia_genes_for_comparison.tsv", sep="\t").head(15)
    top_mepi = pd.read_csv(CMP / "top_mmd_vs_epi_genes_for_comparison.tsv", sep="\t").head(15)
    shared_mia = pd.read_csv(CMP / "shared_genes_aneurysm_vs_mmd_ia.tsv", sep="\t")
    shared_mepi = pd.read_csv(CMP / "shared_genes_aneurysm_vs_mmd_epi.tsv", sep="\t")
    theme_counts = pd.read_csv(CMP / "pathway_theme_keyword_counts.tsv", sep="\t")
    top_shared = pd.read_csv(CMP / "top_shared_genes_effect_matrix.tsv", sep="\t")

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Comparative Transcriptomic Analysis of Aneurysm Rupture and Moya Moya Disease", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Manuscript-style Results and Discussion")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    doc.add_heading("Results", 1)
    doc.add_heading("Cohort-level contrasts and statistical strength", 2)
    doc.add_paragraph(
        f"In the aneurysm cohort (GSE122897), differential expression between ruptured and unruptured tissue yielded "
        f"{summary['comparison_stats']['aneurysm_sig_genes']} genes meeting DESeq2 criteria (FDR < 0.05 and |log2FC| >= 1), "
        f"with robust classifier performance in nested cross-validation (AUC {float(nested.loc[nested['Metric']=='Nested_CV_AUC_Mean','Value'].iloc[0]):.3f} ± "
        f"{float(nested.loc[nested['Metric']=='Nested_CV_AUC_Std','Value'].iloc[0]):.3f})."
    )
    doc.add_paragraph(
        f"In the Moya Moya cohort (GSE157628), the age-proximal contrast (MMD vs IA) produced a substantially weaker signal "
        f"({summary['comparison_stats']['mmd_ia_focus_genes']} focused genes at P < 0.001 and |log2FC| >= 1; no FDR-significant adjusted-model hits), "
        f"whereas the MMD vs EPI contrast showed a much larger transcriptional shift ({summary['comparison_stats']['mmd_epi_focus_genes']} focused genes). "
        f"However, this latter contrast is heavily constrained by cohort imbalance (EPI n=3, all male), and classifier AUC was intentionally suppressed in reporting under reliability rules."
    )

    doc.add_heading("Gene-level overlap and directionality", 2)
    doc.add_paragraph(
        f"Gene overlap analysis showed limited shared signal between aneurysm rupture and MMD vs IA "
        f"(n = {summary['overlap_stats']['a_mia']}), but broader overlap with MMD vs EPI "
        f"(n = {summary['overlap_stats']['a_mepi']}). The three-way overlap across all contrasts was minimal "
        f"(n = {summary['overlap_stats']['a_mia_mepi']}), indicating that disease-context-specific biology predominates."
    )
    doc.add_paragraph(
        f"Shared-gene effect concordance was higher for aneurysm vs MMD-EPI (n = {len(shared_mepi)}) than aneurysm vs MMD-IA (n = {len(shared_mia)}), "
        "but this should be interpreted with caution because the MMD-EPI axis aligns strongly with confound structure."
    )

    add_figure(
        doc,
        FIG / "gene_overlap_3set_ultra_hd.png",
        "Figure 1. Three-set overlap of significant gene signatures across aneurysm and Moya Moya contrasts.",
    )
    total_a = summary["comparison_stats"]["aneurysm_sig_genes"]
    total_mia = summary["comparison_stats"]["mmd_ia_focus_genes"]
    total_mepi = summary["comparison_stats"]["mmd_epi_focus_genes"]
    add_legend_paragraph(
        doc,
        "Legend",
        f"This overlap map quantifies gene-set intersections using the predefined significance cutoffs. "
        f"The aneurysm set contained {total_a} genes, the MMD vs IA set contained {total_mia}, and the MMD vs EPI set contained {total_mepi}. "
        f"Observed pairwise overlaps were aneurysm∩MMDvsIA={summary['overlap_stats']['a_mia']} and aneurysm∩MMDvsEPI={summary['overlap_stats']['a_mepi']}, "
        f"with a three-way overlap of {summary['overlap_stats']['a_mia_mepi']}. "
        f"The disproportionate overlap with MMD vs EPI relative to MMD vs IA reinforces that much of the MMD-EPI contrast likely reflects cohort structure rather than a stable disease-shared core."
    )
    add_legend_paragraph(
        doc,
        "Inference",
        "The overlap architecture supports a selective-convergence model: a robust aneurysm signature, a weak age-proximal MMD-vs-IA signal, "
        "and a broad but potentially confounded MMD-vs-EPI signal. This provides a coherent bridge into effect-direction analyses."
    )
    doc.add_paragraph("")

    add_figure(
        doc,
        FIG / "effect_concordance_scatter_ultra_hd.png",
        "Figure 2. Effect-size concordance for shared gene symbols between aneurysm and Moya Moya contrasts.",
    )
    corr_mia = shared_mia[["log2FoldChange", "Log2FC_case_vs_ctrl"]].corr().iloc[0, 1]
    corr_mepi = shared_mepi[["log2FoldChange", "Log2FC_case_vs_ctrl"]].corr().iloc[0, 1]
    mia_sig_like = int((shared_mia["P_value"] < 0.05).sum())
    mepi_sig_like = int((shared_mepi["FDR_BH"] < 0.05).sum())
    add_legend_paragraph(
        doc,
        "Legend",
        f"Each point is a shared gene symbol, with x-axis effect in aneurysm and y-axis effect in Moya Moya. "
        f"For aneurysm vs MMD-IA, n={len(shared_mia)} shared genes yielded Pearson r={corr_mia:.3f}, and {mia_sig_like} genes had MMD-vs-IA P<0.05. "
        f"For aneurysm vs MMD-EPI, n={len(shared_mepi)} shared genes yielded Pearson r={corr_mepi:.3f}, and {mepi_sig_like} genes had MMD-vs-EPI FDR<0.05. "
        f"Quadrants represent concordant versus discordant directionality."
    )
    add_legend_paragraph(
        doc,
        "Inference",
        "Directionality is not globally conserved, indicating that shared membership does not guarantee shared regulation. "
        "This motivates pathway-level and clustered-pattern interpretation rather than single-axis transfer of biomarker claims."
    )
    doc.add_paragraph("")

    add_figure(
        doc,
        FIG / "shared_gene_clustermap_ultra_hd.png",
        "Figure 3. Clustered effect map of top shared genes across aneurysm, MMD vs IA, and MMD vs EPI.",
    )
    top5 = top_shared.head(5)["GeneSymbol"].tolist()
    add_legend_paragraph(
        doc,
        "Legend",
        f"The clustermap displays the top {len(top_shared)} shared genes ranked by combined absolute effect magnitude across all contrasts. "
        f"Columns are aneurysm, MMD-vs-IA, and MMD-vs-EPI effect sizes transformed to z-scores. "
        f"Highest-ranked genes included {', '.join(top5)}. "
        f"Cluster blocks identify concordant multi-contrast modules versus inversion modules where effect direction flips between contrasts."
    )
    add_legend_paragraph(
        doc,
        "Inference",
        "Gene programs separate into disease-consistent and disease-divergent clusters, supporting a nuanced biology where partial molecular overlap coexists with context-specific rewiring."
    )
    doc.add_paragraph("")

    add_figure(
        doc,
        FIG / "pathway_theme_radar_ultra_hd.png",
        "Figure 4. Comparative pathway-theme footprint across datasets and contrasts.",
    )
    top_a_theme = theme_counts.sort_values("Aneurysm", ascending=False).iloc[0]
    top_mia_theme = theme_counts.sort_values("MMD_vs_IA", ascending=False).iloc[0]
    top_mepi_theme = theme_counts.sort_values("MMD_vs_EPI", ascending=False).iloc[0]
    add_legend_paragraph(
        doc,
        "Legend",
        f"Theme counts summarize keyword-matched enrichment terms by biological domain. "
        f"Aneurysm showed strongest representation in {top_a_theme['Theme']} ({int(top_a_theme['Aneurysm'])} terms), "
        f"MMD-vs-IA in {top_mia_theme['Theme']} ({int(top_mia_theme['MMD_vs_IA'])} terms), "
        f"and MMD-vs-EPI in {top_mepi_theme['Theme']} ({int(top_mepi_theme['MMD_vs_EPI'])} terms). "
        f"Raw counts across themes were: {theme_counts.to_dict(orient='records')}."
    )
    add_legend_paragraph(
        doc,
        "Inference",
        "Despite quantitative differences, all contrasts retain vascular, matrix, and inflammatory components; however, MMD-vs-EPI shows globally inflated theme counts, "
        "consistent with broad transcriptomic separation likely amplified by cohort imbalance."
    )
    doc.add_paragraph("")

    add_table(doc, top_aneurysm, "Table 1. Top aneurysm genes (ruptured vs unruptured)")
    add_legend_paragraph(
        doc,
        "Legend",
        f"Table 1 reports the top {len(top_aneurysm)} aneurysm genes ranked by significance. "
        f"log2FC ranges from {top_aneurysm['Aneurysm_log2FC'].min():.2f} to {top_aneurysm['Aneurysm_log2FC'].max():.2f}; "
        f"adjusted P values range from {top_aneurysm['Aneurysm_padj'].min():.2e} to {top_aneurysm['Aneurysm_padj'].max():.2e}. "
        f"Highest-effect genes include {', '.join(top_aneurysm.sort_values('Aneurysm_log2FC', ascending=False).head(5)['GeneSymbol'].tolist())}."
    )
    add_legend_paragraph(
        doc,
        "Inference",
        "This table anchors the aneurysm-specific signal and serves as the reference axis for cross-condition comparisons."
    )
    doc.add_paragraph("")

    add_table(doc, top_mia, "Table 2. Top Moya Moya genes (MMD vs IA)")
    add_legend_paragraph(
        doc,
        "Legend",
        f"Table 2 reports the top {len(top_mia)} MMD-vs-IA genes. "
        f"log2FC ranges from {top_mia['MMDvsIA_log2FC'].min():.2f} to {top_mia['MMDvsIA_log2FC'].max():.2f}; "
        f"raw P values range from {top_mia['MMDvsIA_P'].min():.2e} to {top_mia['MMDvsIA_P'].max():.2e}; "
        f"FDR values range from {top_mia['MMDvsIA_FDR'].min():.3f} to {top_mia['MMDvsIA_FDR'].max():.3f}."
    )
    add_legend_paragraph(
        doc,
        "Inference",
        "Even among leading genes, multiple-testing control remains weak, reinforcing that this is a lower-power but clinically better-matched contrast."
    )
    doc.add_paragraph("")

    add_table(doc, top_mepi, "Table 3. Top Moya Moya genes (MMD vs EPI)")
    add_legend_paragraph(
        doc,
        "Legend",
        f"Table 3 reports the top {len(top_mepi)} MMD-vs-EPI genes. "
        f"log2FC ranges from {top_mepi['MMDvsEPI_log2FC'].min():.2f} to {top_mepi['MMDvsEPI_log2FC'].max():.2f}; "
        f"raw P values range from {top_mepi['MMDvsEPI_P'].min():.2e} to {top_mepi['MMDvsEPI_P'].max():.2e}; "
        f"FDR values range from {top_mepi['MMDvsEPI_FDR'].min():.2e} to {top_mepi['MMDvsEPI_FDR'].max():.2e}. "
        f"Representative high-magnitude genes include {', '.join(top_mepi.sort_values('MMDvsEPI_log2FC', ascending=False).head(5)['GeneSymbol'].tolist())}."
    )
    add_legend_paragraph(
        doc,
        "Inference",
        "The large effect sizes and broad significance indicate strong separation, but interpretation must remain conservative given demographic imbalance and very small EPI sample size."
    )
    doc.add_paragraph("")

    # =========================================================================
    # NEW SECTION: Shared genes (ruptured vs unruptured aneurysm AND Moya Moya cases vs controls)
    # =========================================================================
    doc.add_heading("Shared Gene Signature: Ruptured vs Unruptured Aneurysms and Moya Moya Cases vs Controls", 1)
    doc.add_paragraph(
        "This section focuses specifically on genes that were differentially expressed in ruptured versus unruptured aneurysm tissue "
        "and that also showed differential expression between Moya Moya cases and controls (either MMD vs IA or MMD vs EPI). "
        "Overlap was defined using strict significance criteria: aneurysm genes (DESeq2 padj < 0.05, |log2FC| >= 1) "
        "intersected with MMD vs IA genes (P < 0.001, |log2FC| >= 1) or MMD vs EPI genes (FDR < 0.05, |log2FC| >= 1). "
        "These overlapping genes represent candidate shared molecular nodes across cerebrovascular pathologies."
    )
    doc.add_paragraph("")

    # Load strict-overlap tables (genes significant in BOTH contrasts)
    strict_mia_path = CMP / "strict_overlap_aneurysm_mmd_ia.tsv"
    strict_mepi_path = CMP / "strict_overlap_aneurysm_mmd_epi.tsv"
    triple_path = CMP / "strict_overlap_triple.tsv"
    if strict_mia_path.exists():
        strict_mia = pd.read_csv(strict_mia_path, sep="\t")
    else:
        strict_mia = pd.DataFrame(columns=["GeneSymbol", "log2FoldChange", "padj", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH"])
    if strict_mepi_path.exists():
        strict_mepi = pd.read_csv(strict_mepi_path, sep="\t")
    else:
        strict_mepi = pd.DataFrame(columns=["GeneSymbol", "log2FoldChange", "padj", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH"])
    if triple_path.exists():
        triple_overlap = pd.read_csv(triple_path, sep="\t")
    else:
        triple_overlap = pd.DataFrame(columns=["GeneSymbol", "log2FoldChange", "Log2FC_MIA", "Log2FC_MEPI"])

    # Subsection: Aneurysm ∩ MMD vs IA (strict overlap)
    doc.add_heading("Shared genes: Aneurysm rupture and MMD vs IA controls", 2)
    if len(strict_mia) > 0:
        strict_mia_display = strict_mia.copy()
        strict_mia_display["Direction_concordance"] = (
            (strict_mia_display["log2FoldChange"] * strict_mia_display["Log2FC_case_vs_ctrl"] > 0)
            .map({True: "Concordant", False: "Discordant"})
        )
        cols = ["GeneSymbol", "log2FoldChange", "padj", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH", "Direction_concordance"]
        cols = [c for c in cols if c in strict_mia_display.columns]
        strict_mia_tbl = strict_mia_display[cols].rename(columns={
            "log2FoldChange": "Aneurysm_Log2FC_RvU",
            "padj": "Aneurysm_padj",
            "Log2FC_case_vs_ctrl": "MMD_vs_IA_Log2FC",
            "P_value": "MMD_vs_IA_P",
            "FDR_BH": "MMD_vs_IA_FDR",
        }).head(25)
        add_table(doc, strict_mia_tbl, "Table 4. Genes differentially expressed in ruptured vs unruptured aneurysms and in MMD vs IA")
        n_conc_mia = int((strict_mia_display["log2FoldChange"] * strict_mia_display["Log2FC_case_vs_ctrl"] > 0).sum())
        n_disc_mia = len(strict_mia_display) - n_conc_mia
        gene_list = ", ".join(strict_mia_tbl["GeneSymbol"].head(10).tolist())
        add_legend_paragraph(
            doc,
            "Legend",
            f"This table lists genes meeting significance criteria in BOTH the aneurysm (ruptured vs unruptured) and MMD vs IA (cases vs controls) contrast. "
            f"Total strict-overlap genes: {len(strict_mia)}. Of these, {n_conc_mia} showed concordant direction (same sign in both contrasts) "
            f"and {n_disc_mia} discordant. Aneurysm log2FC: ruptured minus unruptured; MMD vs IA log2FC: MMD minus IA. "
            f"Genes shown: {gene_list}."
        )
        add_legend_paragraph(
            doc,
            "Inference",
            f"The small overlap (n={len(strict_mia)}) indicates that few genes are robustly differential in both aneurysm rupture and the age-proximal MMD-vs-IA contrast. "
            "Where concordance exists, it may point to common vascular stress or inflammatory modules; discordance suggests context-specific regulation. "
            "These genes warrant targeted validation as potential shared cerebrovascular risk or repair markers."
        )
    else:
        doc.add_paragraph("No genes met the strict significance criteria in both contrasts for the aneurysm ∩ MMD vs IA overlap.")
    doc.add_paragraph("")

    # Subsection: Aneurysm ∩ MMD vs EPI (strict overlap)
    doc.add_heading("Shared genes: Aneurysm rupture and MMD vs EPI controls", 2)
    if len(strict_mepi) > 0:
        strict_mepi_display = strict_mepi.copy()
        strict_mepi_display["Direction_concordance"] = (
            (strict_mepi_display["log2FoldChange"] * strict_mepi_display["Log2FC_case_vs_ctrl"] > 0)
            .map({True: "Concordant", False: "Discordant"})
        )
        cols = ["GeneSymbol", "log2FoldChange", "padj", "Log2FC_case_vs_ctrl", "P_value", "FDR_BH", "Direction_concordance"]
        cols = [c for c in cols if c in strict_mepi_display.columns]
        strict_mepi_tbl = strict_mepi_display[cols].rename(columns={
            "log2FoldChange": "Aneurysm_Log2FC_RvU",
            "padj": "Aneurysm_padj",
            "Log2FC_case_vs_ctrl": "MMD_vs_EPI_Log2FC",
            "P_value": "MMD_vs_EPI_P",
            "FDR_BH": "MMD_vs_EPI_FDR",
        }).head(25)
        add_table(doc, strict_mepi_tbl, "Table 5. Genes differentially expressed in ruptured vs unruptured aneurysms and in MMD vs EPI")
        n_conc_mepi = int((strict_mepi_display["log2FoldChange"] * strict_mepi_display["Log2FC_case_vs_ctrl"] > 0).sum())
        n_disc_mepi = len(strict_mepi_display) - n_conc_mepi
        gene_list = ", ".join(strict_mepi_tbl["GeneSymbol"].head(10).tolist())
        add_legend_paragraph(
            doc,
            "Legend",
            f"This table lists genes meeting significance criteria in BOTH the aneurysm (ruptured vs unruptured) and MMD vs EPI (cases vs controls) contrast. "
            f"Total strict-overlap genes: {len(strict_mepi)}. Of these, {n_conc_mepi} showed concordant direction and {n_disc_mepi} discordant. "
            f"Representative genes include {gene_list}."
        )
        add_legend_paragraph(
            doc,
            "Inference",
            f"The larger overlap (n={len(strict_mepi)}) with MMD vs EPI reflects the broader MMD-EPI transcriptional shift, "
            "but demographic confounding (EPI n=3, all male; large age gap) limits causal interpretation. "
            "Concordant genes may reflect shared inflammatory or vascular remodeling programs; discordant genes highlight disease-specific rewiring. "
            "These findings should be treated as hypothesis-generating and validated in balanced cohorts."
        )
    else:
        doc.add_paragraph("No genes met the strict significance criteria in both contrasts for the aneurysm ∩ MMD vs EPI overlap.")
    doc.add_paragraph("")

    # Subsection: Triple-overlap (all three contrasts, strict)
    doc.add_heading("Triple-overlap genes: Aneurysm, MMD vs IA, and MMD vs EPI", 2)
    if len(triple_overlap) > 0:
        triple_cols = ["GeneSymbol", "log2FoldChange"]
        if "Log2FC_MIA" in triple_overlap.columns:
            triple_cols.append("Log2FC_MIA")
        if "Log2FC_MEPI" in triple_overlap.columns:
            triple_cols.append("Log2FC_MEPI")
        triple_cols = [c for c in triple_cols if c in triple_overlap.columns]
        triple_tbl = triple_overlap[triple_cols].rename(columns={
            "log2FoldChange": "Aneurysm_Log2FC_RvU",
            "Log2FC_MIA": "MMD_vs_IA_Log2FC",
            "Log2FC_MEPI": "MMD_vs_EPI_Log2FC",
        }).head(20)
        add_table(doc, triple_tbl, "Table 6. Genes differentially expressed in ruptured vs unruptured aneurysms and in both MMD vs IA and MMD vs EPI")
        gene_list = ", ".join(triple_tbl["GeneSymbol"].head(8).tolist())
        add_legend_paragraph(
            doc,
            "Legend",
            f"Genes meeting significance criteria in all three contrasts: aneurysm (ruptured vs unruptured), MMD vs IA, and MMD vs EPI. "
            f"Total triple-overlap genes: {len(triple_overlap)}. Genes shown: {gene_list}."
        )
        add_legend_paragraph(
            doc,
            "Inference",
            f"The minimal triple overlap (n={len(triple_overlap)}) underscores that aneurysm rupture and Moya Moya involve largely distinct transcriptional programs. "
            "The few shared genes may represent core vascular or inflammatory nodes amenable to mechanistic follow-up and cross-disease biomarker exploration."
        )
    else:
        doc.add_paragraph("No genes met the strict significance criteria in all three contrasts.")
    doc.add_paragraph("")

    doc.add_heading("Discussion", 1)
    doc.add_paragraph(
        "The comparative analysis indicates that aneurysm rupture biology is characterized by a robust and internally consistent transcriptomic program, "
        "including strong differential expression and coherent pathway enrichment in motility, extracellular matrix, vascular development, and inflammatory signaling domains. "
        "This structure is compatible with a mechanistic rupture cascade centered on endothelial activation, matrix turnover, and immune-cell recruitment."
    )
    doc.add_paragraph(
        "In contrast, the Moya Moya dataset shows a bifurcated profile. The clinically more appropriate MMD vs IA comparison is statistically weak after strict correction, "
        "suggesting either smaller true effects, limited power, or both. The MMD vs EPI contrast appears transcriptionally dramatic but is likely inflated by demographic and cohort effects. "
        "Accordingly, MMD vs EPI should be treated as exploratory and not as primary evidence of disease-specific molecular separation."
    )
    doc.add_paragraph(
        "Cross-condition overlap results support selective convergence rather than global concordance: only a small core of genes is shared across all comparisons, "
        "while most signatures remain context-dependent. This pattern suggests that aneurysm rupture and Moya Moya involve partially intersecting but largely distinct regulatory programs."
    )
    doc.add_paragraph(
        "For manuscript framing, aneurysm findings can be positioned as robust discovery with translational potential, whereas Moya Moya findings are best presented as hypothesis-generating "
        "and in need of balanced, larger-cohort validation with explicit covariate control."
    )

    doc.add_heading("Limitations", 1)
    doc.add_paragraph(
        "Primary limitations include small sample size in Moya Moya subgroup analyses (especially EPI controls), demographic imbalance across contrasts, "
        "and platform differences that can complicate direct cross-disease harmonization. These factors motivate prospective validation and matched-cohort replication."
    )

    out_doc = CMP / "Aneurysm_vs_MoyaMoya_Comparative_Manuscript_Style.docx"
    doc.save(out_doc)
    print(str(out_doc))


if __name__ == "__main__":
    main()
